import importlib
import json
from pathlib import Path
import sys
from types import SimpleNamespace
import pytest

ASSISTANT = Path(__file__).resolve().parents[2] / "car mechanic assistant"
sys.path.insert(0, str(ASSISTANT))
from src.document_processor import DocumentProcessor
from src.embeddings import GeminiEmbeddings
from src.vector_store import VectorStore
from src.rag_system import MechanicAssistant
from web_app import create_app


class FakeEmbeddings:
    model = "test-only-embedding"
    dimensions = 3
    def embed_query(self, text):
        return [1.0, 0.0, 0.0]
    def embed_batch(self, texts, **kwargs):
        return [self.embed_query(text) for text in texts]


@pytest.mark.parametrize("language,text", [("en", "Check battery voltage: 12 V"),
    ("zh", "检查电池电压：12 V"), ("ja", "バッテリー電圧を確認：12 V")])
@pytest.mark.parametrize("filetype", ["pdf", "docx"])
def test_multilingual_pdf_docx(tmp_path, language, text, filetype):
    path = tmp_path / ("manual." + filetype)
    if filetype == "pdf":
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        font = {"en": "Helvetica", "zh": "STSong-Light", "ja": "HeiseiMin-W3"}[language]
        if language != "en":
            pdfmetrics.registerFont(UnicodeCIDFont(font))
        canvas = Canvas(str(path))
        canvas.setFont(font, 12)
        canvas.drawString(40, 700, text)
        canvas.save()
    else:
        from docx import Document
        document = Document()
        document.add_heading("Battery", 1)
        document.add_paragraph(text)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Voltage"
        table.cell(0, 1).text = "12 V"
        document.save(path)
    processor = DocumentProcessor()
    chunks = processor.process_document(path)
    assert text in " ".join(chunk.content for chunk in chunks)
    assert "12 V" in chunks[0].content
    assert chunks[0].metadata["language"] == language
    assert chunks[0].metadata["location"]
    assert processor.manifest["accepted_unique_pdf_pages"] == (1 if filetype == "pdf" else 0)


def test_scan_rejected_and_recorded(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    path = tmp_path / "scan.pdf"
    canvas = Canvas(str(path))
    canvas.rect(40, 600, 100, 100)
    canvas.showPage()
    canvas.save()
    processor = DocumentProcessor()
    with pytest.raises(ValueError, match="needs OCR"):
        processor.process_document(path)
    assert processor.manifest["files"][0]["status"] == "rejected"
    assert processor.manifest["accepted_unique_pdf_pages"] == 0


def test_unique_ids_bounded_chunks_and_metadata(tmp_path):
    for folder in ("a", "b"):
        (tmp_path / folder).mkdir()
        (tmp_path / folder / "same.md").write_text("## Long\n" + "A" * 350)
    chunks = DocumentProcessor(100, 20).process_knowledge_base(tmp_path)
    assert len({chunk.chunk_id for chunk in chunks}) == len(chunks)
    assert all(len(chunk.content) <= 100 for chunk in chunks)
    processor = DocumentProcessor()
    tags = processor._sanitize_metadata({"manufacturer": " BmW ", "nothing": None, "year_range": "2007-2011", "codes": ["P0300"]})
    assert tags["make"] == "bmw" and "nothing" not in tags and tags["year_start"] == 2007
    with pytest.raises(ValueError):
        processor._sanitize_metadata({"make": ["BMW", "Honda"]})
    with pytest.raises(ValueError):
        DocumentProcessor(100, 100)


def test_embedding_failure_never_becomes_zero_vector():
    class Models:
        def embed_content(self, **kwargs):
            raise TimeoutError("provider unavailable")
    embeddings = GeminiEmbeddings(client=SimpleNamespace(models=Models()), dimensions=3)
    with pytest.raises(TimeoutError):
        embeddings.embed_batch(["A", "B"])
    embeddings.client.models.embed_content = lambda **kwargs: SimpleNamespace(embeddings=[SimpleNamespace(values=[0, 0, 0])])
    with pytest.raises(ValueError, match="zero norm"):
        embeddings.embed_query("test")


def make_corpus(tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    for name, metadata, text in [
        ("generic", {"scope": "generic"}, "P0300 misfire diagnostic checks"),
        ("bmw", {"make": " BmW ", "year_range": "2007-2011"}, "BMW P0300 checks"),
        ("honda", {"make": "Honda", "year_range": "2007-2011"}, "Honda P0300 checks")]:
        (kb / (name + ".md")).write_text(text + "\n```json\n" + json.dumps(metadata) + "\n```")
    return DocumentProcessor().process_knowledge_base(kb)


def test_real_chroma_filters_keep_generic_and_exclude_wrong_make(tmp_path):
    chunks = make_corpus(tmp_path)
    store = VectorStore(str(tmp_path / "index"), embeddings=FakeEmbeddings())
    store.add_documents(chunks)
    store.add_documents(chunks)
    assert store.collection.count() == 3
    for make in ("BMW", "bmw", " BmW "):
        results = store.query_with_reranking("P0300", vehicle_info={"make": make, "year": 2007})
        files = {result["metadata"]["filename"] for result in results}
        assert files == {"bmw.md", "generic.md"}
    results = store.query_with_reranking("P0300", vehicle_info={"make": "Toyota"})
    assert [result["metadata"]["filename"] for result in results] == ["generic.md"]
    assert not store.query_with_reranking("P0300", vehicle_info={"make": "BMW", "year": 2020})[0]["metadata"].get("make")


def test_index_replacement_and_invalid_batch(tmp_path):
    chunks = make_corpus(tmp_path)
    store = VectorStore(str(tmp_path / "index"), embeddings=FakeEmbeddings())
    store.add_documents(chunks)
    path = tmp_path / "kb" / "bmw.md"
    path.write_text('Updated P0300 checks\n```json\n{"make": "BMW"}\n```')
    updated = DocumentProcessor().process_document(path, tmp_path / "kb")
    store.add_documents(updated)
    assert store.collection.count() == 3
    assert chunks[0].chunk_id not in store.collection.get()["ids"]
    class Broken(FakeEmbeddings):
        def embed_batch(self, texts, **kwargs):
            return [[0, 0, 0] for _ in texts]
    store.embeddings = Broken()
    with pytest.raises(ValueError):
        store.add_documents(updated)
    assert store.collection.count() == 3


def test_old_index_untouched_and_identity_checked(tmp_path):
    path = tmp_path / "legacy"
    path.mkdir()
    (path / "chroma.sqlite3").write_bytes(b"original")
    with pytest.raises(ValueError, match="Legacy"):
        VectorStore(str(path), embeddings=FakeEmbeddings())
    assert (path / "chroma.sqlite3").read_bytes() == b"original"
    path = tmp_path / "new"
    VectorStore(str(path), embeddings=FakeEmbeddings())
    class Different(FakeEmbeddings):
        model = "another-model"
    with pytest.raises(ValueError, match="mismatch"):
        VectorStore(str(path), embeddings=Different())


def test_vehicle_regex_and_all_policy_text():
    assistant = MechanicAssistant.__new__(MechanicAssistant)
    info = assistant.extract_vehicle_info("2007 bmw E90 n52")
    assert info == {"make": "bmw", "year": 2007, "model": "e90", "engine": "n52"}
    assert "make" not in assistant.extract_vehicle_info("know how to replace plugs")
    class Store:
        def get_workshop_rules(self):
            return "Rule A\nRule B"
    assistant.vector_store = Store()
    context = assistant.format_context([])
    assert "Rule A" in context and "Rule B" in context


def test_web_input_feedback_and_error_redaction(tmp_path):
    class Assistant:
        def query(self, text):
            if text == "fail":
                raise ValueError("secret=private")
            return {"query": text, "response": "Answer", "sources": [], "num_sources": 0}
    client = create_app(Assistant(), tmp_path / "feedback.sqlite").test_client()
    assert client.get("/health").status_code == 200
    assert client.post("/query", json={"query": 123}).status_code == 400
    assert client.post("/query", json={"query": "A" * 4001}).status_code == 400
    result = client.post("/query", json={"query": "P0300"}).json
    assert client.post("/feedback", json={"query_id": result["query_id"], "rating": "useful", "comment": "Source relevant"}).status_code == 200
    assert client.post("/feedback", json={"query_id": "missing", "rating": "useful"}).status_code == 404
    assert client.post("/feedback", json={"query_id": result["query_id"], "rating": "invented"}).status_code == 400
    response = client.post("/query", json={"query": "fail"})
    assert response.status_code == 503 and "private" not in response.text


def test_evaluation_does_not_confuse_hit_with_precision():
    from src.evaluate_retrieval import evaluate
    labels = [{"query_id": "q1", "category": "misfire_ignition", "language": "en", "relevant_ids": ["correct"],
               "reviewer": "fixture reviewer", "reviewed_at": "2026-10-09"}]
    report = evaluate(labels, [{"query_id": "q1", "retrieved_ids": ["correct", "correct", "wrong"]}])
    assert report["precision_at_5"] == 0.2
    assert report["hit_at_5"] == 1
    assert not report["all_eight_categories_present"]
    assert evaluate(labels, [])["precision_at_5"] == 0
    with pytest.raises(ValueError):
        evaluate(labels, [{"query_id": "unknown", "retrieved_ids": []}])


def test_embedding_pacing_and_bounded_retry(monkeypatch):
    from src import embeddings as module
    from google.genai.errors import ClientError
    sleeps = []
    monkeypatch.setattr(module.time, "sleep", sleeps.append)
    calls = []
    def embed(**kwargs):
        calls.append(kwargs)
        if len(calls) == 1:
            raise ClientError(429, {"error": {"message": "rate limited"}})
        return SimpleNamespace(embeddings=[SimpleNamespace(values=[1, 0, 0]) for _ in kwargs["contents"]])
    embeddings = GeminiEmbeddings(client=SimpleNamespace(models=SimpleNamespace(embed_content=embed)), dimensions=3)
    assert len(embeddings.embed_batch(["text"] * 33)) == 33
    assert sleeps == [2, 21]
    assert len(calls) == 3
    calls.clear()
    def unavailable(**kwargs):
        calls.append(True)
        raise ClientError(401, {"error": {"message": "unauthorised"}})
    embeddings.client.models.embed_content = unavailable
    with pytest.raises(ClientError):
        embeddings.embed_batch(["text"])
    assert len(calls) == 1


def test_generation_retry_and_source_payload(monkeypatch):
    from src import rag_system as module
    from google.genai.errors import ServerError
    sleeps, calls = [], []
    monkeypatch.setattr(module.time, "sleep", sleeps.append)
    assistant = MechanicAssistant.__new__(MechanicAssistant)
    assistant.model_name = "test-generation"
    assistant.generation_config = None
    def generate(**kwargs):
        calls.append(kwargs)
        if len(calls) == 1:
            raise ServerError(503, {"error": {"message": "unavailable"}})
        return SimpleNamespace(text="Check voltage [passage-1, page 2].")
    assistant.client = SimpleNamespace(models=SimpleNamespace(generate_content=generate))
    assert assistant._generate("question").text.startswith("Check")
    assert len(calls) == 2 and sleeps == [2]
    assistant.client.models.generate_content = lambda **kwargs: SimpleNamespace(text=None)
    with pytest.raises(ValueError, match="no answer"):
        assistant._generate("question")
    passage = {"id": "passage-1", "content": "Check voltage", "relevance_score": 1,
               "metadata": {"folder": "manuals", "filename": "battery.pdf", "location": "page 2",
                            "source_url": "https://example.com/manual", "corpus_kind": "public_reference_placeholder", "language": "en"}}
    assistant.vector_store = SimpleNamespace(embeddings=FakeEmbeddings(), identity={"schema": 2},
                                           get_workshop_rules=lambda: "Policy")
    assistant.retrieve_context = lambda *args, **kwargs: [passage]
    assistant.generate_response = lambda *args: "Answer [source:passage-1]"
    result = assistant.query("question")
    assert result["sources"][0]["id"] == "passage-1"
    assert result["sources"][0]["location"] == "page 2"
    assert result["sources"][0]["excerpt"] == "Check voltage"
    assert result["sources"][0]["source_url"] == "https://example.com/manual"
    assert result["sources"][0]["corpus_kind"] == "public_reference_placeholder"
    assert result["sources"][0]["language"] == "en"
    assert result["metadata"]["index_identity"] == {"schema": 2}
